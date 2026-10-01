"""Renders resume.json to JoseLoboPortfolioResume.docx (repo root).
Manual tool, not part of the site build:
    python3 -m pip install python-docx && python3 scripts/resume/build-docx.py
"""
import json
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches

here = Path(__file__).parent
r = json.loads((here / "resume.json").read_text())
BLUE = RGBColor(0x2B, 0x5B, 0xD7)

doc = Document()
for s in doc.sections:
    s.top_margin = s.bottom_margin = Inches(0.6)
    s.left_margin = s.right_margin = Inches(0.7)
normal = doc.styles["Normal"]
normal.font.name = "Arial"
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(1)


def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    run = p.add_run(text.upper())
    run.bold = True
    run.font.color.rgb = BLUE
    run.font.size = Pt(11)


def bullets(items):
    for i in items:
        doc.add_paragraph(i, style="List Bullet")


title = doc.add_paragraph()
t = title.add_run(r["name"])
t.bold = True
t.font.size = Pt(22)
doc.add_paragraph("  |  ".join(r["contact"]))

heading("Professional Summary")
doc.add_paragraph(r["summary"])

heading("Professional Experience")
for j in r["experience"]:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.add_run(f'{j["org"]} · {j["role"]}').bold = True
    p.add_run(f'    {j["dates"]}').italic = True
    if j.get("note"):
        n = doc.add_paragraph().add_run(j["note"])
        n.italic = True
        n.font.size = Pt(9)
    bullets(j["bullets"])

heading("Technical Skills")
for k, v in r["skills"]:
    p = doc.add_paragraph()
    p.add_run(f"{k}: ").bold = True
    p.add_run(v)

heading("Education")
for k, v in r["education"]:
    p = doc.add_paragraph()
    p.add_run(k).bold = True
    p.add_run(f" · {v}")

heading("Certifications")
bullets(r["certifications"])
heading("Clinical Experience")
bullets(r["clinical"])
heading("Volunteer Experience")
bullets(r["volunteer"])
heading("Presentations and Lectures")
bullets(r["presentations"])

doc.save(here / "../../JoseLoboPortfolioResume.docx")
print("wrote JoseLoboPortfolioResume.docx")
